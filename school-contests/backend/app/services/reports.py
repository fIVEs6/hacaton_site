"""Отчёты: сбор данных и выгрузка в XLSX / CSV / PDF."""
import csv
import html
import io
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import or_, select
from sqlalchemy.orm import Session
from weasyprint import HTML

from app.models import (LEVEL_RU, STATUS_RU, Diploma, Event, ParticipationType, Registration, Result, Team, User,
                        UserRole)

KINDS = {"students": "Отчёт по ученикам", "classes": "Отчёт по классам",
         "events": "Отчёт по мероприятиям", "period": "Отчёт за период"}


def _participations(db: Session, grade: int | None, event_id: int | None,
                    date_from: date | None, date_to: date | None) -> list[dict]:
    """Одна строка = одно участие ученика в мероприятии вместе с результатом и дипломом."""
    stmt = select(Registration, Event, User).join(Event, Event.id == Registration.event_id) \
        .join(User, User.id == Registration.user_id)
    if grade:
        stmt = stmt.where(User.grade == grade)
    if event_id:
        stmt = stmt.where(Event.id == event_id)
    if date_from:
        stmt = stmt.where(Event.end_date >= date_from)
    if date_to:
        stmt = stmt.where(Event.start_date <= date_to)
    rows = db.execute(stmt.order_by(Event.start_date, User.last_name)).all()
    ev_ids = {e.id for _, e, _ in rows}
    by_reg, by_team = {}, {}
    for r in db.scalars(select(Result).where(Result.event_id.in_(list(ev_ids)))) if ev_ids else []:
        if r.registration_id:
            by_reg[r.registration_id] = r
        if r.team_id:
            by_team[r.team_id] = r
    diplomas = {(d.event_id, d.user_id): d.number for d in db.scalars(
        select(Diploma).where(Diploma.revoked_at.is_(None)))}
    teams = {t.id: t.name for t in db.scalars(select(Team))}
    out = []
    for reg, ev, user in rows:
        res = by_reg.get(reg.id) or (by_team.get(reg.team_id) if reg.team_id else None)
        out.append({"reg": reg, "event": ev, "user": user, "result": res,
                    "team": teams.get(reg.team_id), "diploma": diplomas.get((ev.id, user.id))})
    return out


def build(db: Session, kind: str, grade: int | None = None, event_id: int | None = None,
          date_from: date | None = None, date_to: date | None = None) -> tuple[str, list[str], list[list]]:
    title = KINDS[kind]
    if kind == "events":
        events = db.scalars(select(Event).order_by(Event.start_date.desc())).all()
        parts = _participations(db, None, None, date_from, date_to)
        agg = defaultdict(lambda: [0, 0, 0, 0])
        for p in parts:
            a = agg[p["event"].id]
            a[0] += 1
            a[1] += p["reg"].status.value == "prizewinner"
            a[2] += p["reg"].status.value == "winner"
            a[3] += bool(p["diploma"])
        headers = ["Мероприятие", "Предмет", "Уровень", "Формат", "Даты", "Статус", "Участников",
                   "Призёров", "Победителей", "Дипломов"]
        status = {"draft": "Черновик", "published": "Опубликовано", "finished": "Завершено"}
        rows = []
        for e in events:
            if date_from and e.end_date < date_from or date_to and e.start_date > date_to:
                continue
            a = agg[e.id]
            rows.append([e.title, e.subject, LEVEL_RU[e.level.value],
                         "Командный" if e.participation_type == ParticipationType.team else "Индивидуальный",
                         f"{e.start_date:%d.%m.%Y} — {e.end_date:%d.%m.%Y}", status[e.status.value], *a])
        return title, headers, rows

    parts = _participations(db, grade, event_id, date_from, date_to)

    if kind == "students":
        agg: dict[int, dict] = {}
        students = db.scalars(select(User).where(User.role == UserRole.student)
                              .where(User.grade == grade if grade else True)
                              .order_by(User.grade, User.class_letter, User.last_name)).all()
        for u in students:
            agg[u.id] = {"user": u, "n": 0, "part": 0, "prize": 0, "win": 0, "score": 0.0, "dip": 0}
        for p in parts:
            a = agg.get(p["user"].id)
            if not a:
                continue
            a["n"] += 1
            st = p["reg"].status.value
            a["part"] += st == "participated"
            a["prize"] += st == "prizewinner"
            a["win"] += st == "winner"
            a["score"] += float(p["result"].score) if p["result"] and p["result"].score is not None else 0
            a["dip"] += bool(p["diploma"])
        headers = ["ФИО", "Класс", "Email", "Участий", "Просто участник", "Призёр", "Победитель", "Сумма баллов", "Дипломов"]
        rows = [[a["user"].full_name, a["user"].class_name, a["user"].email, a["n"], a["part"], a["prize"],
                 a["win"], round(a["score"], 2), a["dip"]] for a in agg.values()]
        return title, headers, rows

    if kind == "classes":
        agg = defaultdict(lambda: {"students": set(), "n": 0, "prize": 0, "win": 0, "score": 0.0})
        for p in parts:
            a = agg[p["user"].class_name or "—"]
            a["students"].add(p["user"].id)
            a["n"] += 1
            a["prize"] += p["reg"].status.value == "prizewinner"
            a["win"] += p["reg"].status.value == "winner"
            a["score"] += float(p["result"].score) if p["result"] and p["result"].score is not None else 0
        headers = ["Класс", "Учеников участвовало", "Участий", "Призёров", "Победителей", "Сумма баллов"]

        def sort_key(k):
            digits = "".join(c for c in k if c.isdigit())
            return (int(digits) if digits else 99, k)
        rows = [[k, len(a["students"]), a["n"], a["prize"], a["win"], round(a["score"], 2)]
                for k, a in sorted(agg.items(), key=lambda kv: sort_key(kv[0]))]
        return title, headers, rows

    # period: подробный список участий
    headers = ["Дата", "Мероприятие", "ФИО", "Класс", "Команда", "Статус", "Место", "Баллы", "Диплом №"]
    rows = []
    for p in parts:
        r = p["result"]
        rows.append([p["event"].start_date.strftime("%d.%m.%Y"), p["event"].title, p["user"].full_name,
                     p["user"].class_name, p["team"] or "", STATUS_RU[p["reg"].status.value],
                     r.place if r and r.place else "", float(r.score) if r and r.score is not None else "",
                     p["diploma"] or ""])
    if date_from or date_to:
        title += f" ({date_from.strftime('%d.%m.%Y') if date_from else '…'} — {date_to.strftime('%d.%m.%Y') if date_to else '…'})"
    return title, headers, rows


# ---------- Форматы вывода ----------

def to_xlsx(title: str, headers: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(headers)
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2A44")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200]]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 8), 55)
    ws.freeze_panes = "A4"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(headers: list[str], rows: list[list]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(headers)
    w.writerows(rows)
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM, чтобы Excel открыл кириллицу правильно


def to_pdf(title: str, headers: list[str], rows: list[list]) -> bytes:
    head = "".join(f"<th>{html.escape(h)}</th>" for h in headers)
    body = "".join("<tr>" + "".join(f"<td>{html.escape(str(c))}</td>" for c in r) + "</tr>" for r in rows)
    doc = f"""<html><head><meta charset="utf-8"><style>
      @page {{ size: A4 landscape; margin: 14mm; @bottom-right {{ content: counter(page) " / " counter(pages); font-size: 8pt; }} }}
      body {{ font-family: "Liberation Sans", "DejaVu Sans", sans-serif; font-size: 8.5pt; color: #1f2a44; }}
      h1 {{ font-family: "Liberation Serif", serif; font-size: 16pt; margin: 0 0 4mm 0; }}
      table {{ border-collapse: collapse; width: 100%; }}
      th {{ background: #1f2a44; color: #fff; text-align: left; padding: 2mm; }}
      td {{ border-bottom: 0.2mm solid #c9c2b2; padding: 1.5mm 2mm; }}
      tr:nth-child(even) td {{ background: #f6f0e4; }}
    </style></head><body><h1>{html.escape(title)}</h1><p>Сформирован: {date.today():%d.%m.%Y}. Строк: {len(rows)}</p>
    <table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></body></html>"""
    return HTML(string=doc).write_pdf()
